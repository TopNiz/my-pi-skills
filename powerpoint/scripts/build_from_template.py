#!/usr/bin/env python3
"""Build a new editable deck from an existing PowerPoint template.

This helper is for creating a derived deck from a reference template. It keeps
the template masters, layouts, theme, and placeholder styles, removes the
template's sample slides only when --clear-slides is supplied, and populates
new slides from a JSON specification.

Reference templates are not part of this skill and are never committed. They
live outside the repository, in the directory configured by
``POWERPOINT_TEMPLATES_DIR`` in the skill-local ``.env`` (see ``.env.example``).
--template therefore accepts either an explicit path or a bare filename that is
resolved inside that directory.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATES_DIR_ENV = "POWERPOINT_TEMPLATES_DIR"


def load_skill_env(skill_dir: Path = SKILL_DIR) -> dict[str, str]:
    """Read the skill-local ``.env`` (see ``.env.example``).

    Dependency-free on purpose: the skill must run with nothing but
    ``python-pptx`` installed. Values are returned rather than exported, so
    nothing leaks into the process environment.
    """
    values: dict[str, str] = {}
    env_file = skill_dir / ".env"
    if not env_file.is_file():
        return values
    for raw_line in env_file.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def templates_dir(cli_value: str | None = None) -> Path | None:
    """Resolve the configured template-library directory, if any.

    Precedence: ``--templates-dir``, then the process environment, then the
    skill-local ``.env``.
    """
    raw = (
        cli_value
        or os.environ.get(TEMPLATES_DIR_ENV)
        or load_skill_env().get(TEMPLATES_DIR_ENV)
    )
    return Path(raw).expanduser() if raw else None


def resolve_template(value: str, directory: Path | None) -> Path:
    """Accept an explicit path, or a bare filename inside the configured dir."""
    candidate = Path(value).expanduser()
    if candidate.is_file():
        return candidate
    if directory is not None and not candidate.is_absolute():
        in_directory = directory / candidate
        if in_directory.is_file():
            return in_directory
    if directory is None:
        raise FileNotFoundError(
            f"template {value!r} not found. Pass a path, or configure "
            f"{TEMPLATES_DIR_ENV} in {SKILL_DIR / '.env'} (see .env.example)."
        )
    raise FileNotFoundError(
        f"template {value!r} not found, and it is not in {directory}. "
        f"Check {TEMPLATES_DIR_ENV} in {SKILL_DIR / '.env'}."
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--template",
        required=True,
        help="Template deck: an explicit path, or a filename inside the "
             f"configured ${TEMPLATES_DIR_ENV} directory",
    )
    parser.add_argument(
        "--templates-dir",
        default=None,
        help=f"Override ${TEMPLATES_DIR_ENV} for this run",
    )
    parser.add_argument("--spec", required=True, type=Path,
                        help="JSON slide specification; see references/build-from-template.md")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--clear-slides", action="store_true",
                        help="Remove existing sample slides before adding specified slides")
    parser.add_argument("--list-layouts", action="store_true",
                        help="Print layout names and placeholder metadata, then exit")
    return parser.parse_args()


def layout_by_name(prs: Presentation, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    available = ", ".join(layout.name for layout in prs.slide_layouts)
    raise KeyError(f"layout {name!r} not found; available layouts: {available}")


def placeholder_by_idx(slide, idx: int):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    raise KeyError(f"placeholder idx={idx} not found on layout {slide.slide_layout.name!r}")


def put_text(shape, text: str) -> None:
    if not shape.has_text_frame:
        raise TypeError(f"shape {shape.name!r} has no text frame")
    frame = shape.text_frame
    frame.clear()
    for line_number, line in enumerate(str(text).split("\n")):
        paragraph = frame.paragraphs[0] if line_number == 0 else frame.add_paragraph()
        paragraph.text = line


def put_list(shape, items: list[str]) -> None:
    if not shape.has_text_frame:
        raise TypeError(f"shape {shape.name!r} has no text frame")
    frame = shape.text_frame
    frame.clear()
    for item_number, item in enumerate(items):
        paragraph = frame.paragraphs[0] if item_number == 0 else frame.add_paragraph()
        paragraph.text = str(item)


def clear_slides(prs: Presentation) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        relationship_id = slide_id.get(R_NS)
        prs.part.drop_rel(relationship_id)
        slide_ids.remove(slide_id)


def print_layouts(prs: Presentation) -> None:
    for index, layout in enumerate(prs.slide_layouts):
        print(f"{index}: {layout.name}")
        for placeholder in layout.placeholders:
            fmt = placeholder.placeholder_format
            print(
                f"  idx={fmt.idx} type={fmt.type} name={placeholder.name} "
                f"box=({placeholder.left},{placeholder.top},{placeholder.width},{placeholder.height})"
            )


def populate_slide(prs: Presentation, slide_spec: dict) -> None:
    slide = prs.slides.add_slide(layout_by_name(prs, slide_spec["layout"]))
    title = slide_spec.get("title")
    if title is not None:
        if slide.shapes.title is None:
            raise KeyError(f"layout {slide.slide_layout.name!r} has no title placeholder")
        put_text(slide.shapes.title, title)

    for raw_idx, text in slide_spec.get("text", {}).items():
        put_text(placeholder_by_idx(slide, int(raw_idx)), text)
    for raw_idx, items in slide_spec.get("lists", {}).items():
        put_list(placeholder_by_idx(slide, int(raw_idx)), items)
    for raw_idx, image_path in slide_spec.get("images", {}).items():
        placeholder_by_idx(slide, int(raw_idx)).insert_picture(str(Path(image_path)))


def validate(path: Path, expected_slides: int) -> None:
    check = Presentation(str(path))
    if len(check.slides) != expected_slides:
        raise AssertionError(f"expected {expected_slides} slides, found {len(check.slides)}")
    for slide in check.slides:
        if slide.slide_layout is None:
            raise AssertionError("slide has no associated layout")
    with zipfile.ZipFile(path) as package:
        bad_entry = package.testzip()
        if bad_entry is not None:
            raise AssertionError(f"corrupt package entry: {bad_entry}")


def main() -> int:
    args = parse_args()
    template_path = resolve_template(args.template, templates_dir(args.templates_dir))
    prs = Presentation(str(template_path))
    if args.list_layouts:
        print_layouts(prs)
        return 0

    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("spec must contain a non-empty 'slides' list")
    if args.clear_slides:
        clear_slides(prs)
    for slide_spec in slides:
        populate_slide(prs, slide_spec)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.output))
    validate(args.output, len(slides) if args.clear_slides else len(prs.slides))
    print(f"created: {args.output}")
    print(f"slides: {len(prs.slides)}")
    print("validation: reopened successfully; ZIP package is valid")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (KeyError, TypeError, ValueError, AssertionError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(2)
